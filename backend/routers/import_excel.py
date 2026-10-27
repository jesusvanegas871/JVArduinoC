from fastapi import APIRouter, Depends, HTTPException, UploadFile, File, Form, status
from sqlalchemy.orm import Session
from typing import List, Optional
import openpyxl
import io
import json
from datetime import datetime, timezone

from backend.database import get_db
from backend.models import Product, Category, Location, StockLocation, InventoryMovement, SalesHistory, ExcelImportLog
from backend.schemas import (
    ImportPreviewResponse, ImportRowPreview, ImportConfirmRequest, ImportResult,
    CatalogMetrics, ExcelImportLogOut
)

router = APIRouter(prefix="/api/import", tags=["import"])

def calculate_catalog_metrics(db: Session, new_imported_count: int = 0, duplicates_in_import: int = 0) -> CatalogMetrics:
    total_existing = db.query(Product).count()
    total_active_eligible = db.query(Product).filter(
        Product.is_active == True,
        Product.annual_physical_count < 5
    ).count()

    total_unique_imported = total_existing
    target = 600
    missing = max(0, target - total_existing)

    return CatalogMetrics(
        target_catalog=target,
        total_unique_imported=total_unique_imported,
        total_existing_in_db=total_existing,
        total_new_imported=new_imported_count,
        total_duplicates_detected=duplicates_in_import,
        total_missing_for_600=missing,
        total_active_eligible_for_count=total_active_eligible
    )

@router.get("/metrics", response_model=CatalogMetrics)
def get_catalog_metrics(db: Session = Depends(get_db)):
    return calculate_catalog_metrics(db)

@router.post("/preview", response_model=ImportPreviewResponse)
async def preview_excel_import(
    file: UploadFile = File(...),
    sheet_name: str = Form(...),
    db: Session = Depends(get_db)
):
    if not file.filename.endswith(('.xlsx', '.xls')):
        raise HTTPException(status_code=400, detail="El archivo debe ser un documento Excel (.xlsx)")

    contents = await file.read()
    try:
        wb = openpyxl.load_workbook(io.BytesIO(contents), data_only=True)
    except Exception as e:
        raise HTTPException(status_code=400, detail=f"No se pudo leer el archivo Excel: {str(e)}")

    if sheet_name not in wb.sheetnames:
        raise HTTPException(status_code=400, detail=f"La hoja '{sheet_name}' no existe en el archivo. Hojas disponibles: {', '.join(wb.sheetnames)}")

    sheet = wb[sheet_name]
    rows = list(sheet.iter_rows(values_only=True))

    if not rows:
        raise HTTPException(status_code=400, detail="La hoja de Excel está vacía.")

    headers = []
    header_row_idx = 0

    # Locate header row
    for idx, row in enumerate(rows[:10]):
        row_str = [str(cell).strip().upper() for cell in row if cell is not None]
        if any("CODIGO" in cell or "CÓDIGO" in cell or "REFERENCIA" in cell for cell in row_str):
            header_row_idx = idx
            headers = [str(c).strip() if c is not None else "" for c in row]
            break

    if not headers:
        header_row_idx = 0
        headers = [str(c).strip() if c is not None else "" for c in rows[0]]

    # Map headers to indices
    header_indices = {}
    for i, h in enumerate(headers):
        h_norm = h.upper().strip()
        header_indices[h_norm] = i

    # Existing products in DB for comparison
    existing_products_map = {p.sku.upper(): p for p in db.query(Product).all()}

    preview_rows = []
    seen_file_skus = set()
    file_duplicates_count = 0

    to_create = 0
    to_update = 0
    skipped = 0
    errors = 0

    # Process rows
    for r_idx, row in enumerate(rows[header_row_idx + 1:], start=header_row_idx + 2):
        if not row or all(v is None or str(v).strip() == "" for v in row):
            continue

        sku = ""
        name = ""
        category_name = None
        warnings = []
        row_errors = []
        aux_data = {}

        if "PRODUCTOS A INCLUIR" in sheet_name.upper():
            # Standard columns: CODIGO, DESCRIPCIÓN, CATEGORIA, ENERO ... DICIEMBRE
            code_col = header_indices.get("CODIGO", 0)
            desc_col = header_indices.get("DESCRIPCIÓN", 1) if "DESCRIPCIÓN" in header_indices else header_indices.get("DESCRIPCION", 1)
            cat_col = header_indices.get("CATEGORIA", 2) if "CATEGORIA" in header_indices else header_indices.get("CATEGORÍA", 2)

            raw_code = row[code_col] if code_col < len(row) else None
            raw_desc = row[desc_col] if desc_col < len(row) else None
            raw_cat = row[cat_col] if cat_col < len(row) else None

            if raw_code is not None:
                sku = str(raw_code).strip()
                if sku.endswith('.0'):
                    sku = sku[:-2]
            if raw_desc is not None:
                name = str(raw_desc).strip()
            if raw_cat is not None:
                category_name = str(raw_cat).strip()

            # Aux monthly projections
            months = ["ENERO", "FEBRERO", "MARZO", "ABRIL", "MAYO", "JUNIO", "JULIO", "AGOSTO", "SEPTIEMBRE", "OCTUBRE", "NOVIEMBRE", "DICIEMBRE"]
            monthly_dict = {}
            for m in months:
                if m in header_indices and header_indices[m] < len(row):
                    val = row[header_indices[m]]
                    if val is not None and str(val).strip() != "":
                        try:
                            monthly_dict[m] = int(float(val))
                        except ValueError:
                            monthly_dict[m] = 0
                    else:
                        monthly_dict[m] = None
            aux_data["monthly_projections"] = monthly_dict

        else:
            # DATA VENTAS or generic
            # Columns: Codigo de la referencia, Descripción de la referencia, Fecha, Cantidad por referencia, Empresa...
            code_col = header_indices.get("CODIGO DE LA REFERENCIA", header_indices.get("CODIGO", 4))
            desc_col = header_indices.get("DESCRIPCIÓN DE LA REFERENCIA", header_indices.get("DESCRIPCION DE LA REFERENCIA", header_indices.get("DESCRIPCIÓN", 5)))
            cat_col = header_indices.get("CALEGORIA", header_indices.get("CATEGORIA", 10))

            raw_code = row[code_col] if code_col < len(row) else None
            raw_desc = row[desc_col] if desc_col < len(row) else None
            raw_cat = row[cat_col] if cat_col < len(row) and row[cat_col] and not str(row[cat_col]).startswith("=") else None

            if raw_code is not None:
                sku = str(raw_code).strip()
                if sku.endswith('.0'):
                    sku = sku[:-2]
            if raw_desc is not None:
                name = str(raw_desc).strip()
            if raw_cat is not None:
                category_name = str(raw_cat).strip()

            # Additional sales record fields
            order_num_col = header_indices.get("NO. DE ORDEN", 0)
            date_col = header_indices.get("FECHA", 1)
            qty_col = header_indices.get("CANTIDAD POR REFERENCIA", 6)
            total_val_col = header_indices.get("VALOR TOTAL DE LAS REFERENCIAS", 7)
            company_col = header_indices.get("EMPRESA", 8)
            city_col = header_indices.get("CIUDAD", 9)

            aux_data["order_number"] = str(row[order_num_col]).strip() if order_num_col < len(row) and row[order_num_col] is not None else None
            aux_data["sale_date"] = str(row[date_col]) if date_col < len(row) and row[date_col] is not None else None
            aux_data["quantity"] = int(row[qty_col]) if qty_col < len(row) and row[qty_col] is not None and str(row[qty_col]).isdigit() else 1
            aux_data["company"] = str(row[company_col]).strip() if company_col < len(row) and row[company_col] is not None else None
            aux_data["city"] = str(row[city_col]).strip() if city_col < len(row) and row[city_col] is not None else None

        # Validation
        if not sku:
            row_errors.append("El código de referencia es obligatorio.")
        if not name:
            row_errors.append("La descripción o nombre de referencia es obligatoria.")

        sku_upper = sku.upper()
        row_status = "NUEVO"
        action = "CREAR"

        if row_errors:
            row_status = "INVALIDO"
            action = "ERROR"
            errors += 1
        elif sku_upper in seen_file_skus:
            row_status = "DUPLICADO_EN_ARCHIVO"
            file_duplicates_count += 1
            if "DATA VENTAS" in sheet_name.upper():
                action = "REGISTRO_VENTA_AUX"
                row_status = "VENTA_HISTORICA"
            else:
                action = "OMITIR"
                warnings.append(f"El código '{sku}' aparece múltiples veces en este archivo. Se conservará la primera ocurrencia.")
                skipped += 1
        else:
            seen_file_skus.add(sku_upper)
            if sku_upper in existing_products_map:
                existing_p = existing_products_map[sku_upper]
                row_status = "EXISTENTE"
                action = "ACTUALIZAR"
                to_update += 1
                if existing_p.name.strip().lower() != name.strip().lower():
                    warnings.append(f"La descripción en archivo ('{name}') difiere de la registrada en catálogo ('{existing_p.name}').")
            else:
                to_create += 1

        preview_rows.append(ImportRowPreview(
            row_index=r_idx,
            sku=sku,
            name=name,
            category=category_name,
            origin_sheet=sheet_name,
            status=row_status,
            action=action,
            warnings=warnings,
            errors=row_errors,
            auxiliary_data=aux_data
        ))

    metrics = calculate_catalog_metrics(
        db,
        new_imported_count=to_create,
        duplicates_in_import=file_duplicates_count
    )

    return ImportPreviewResponse(
        filename=file.filename,
        sheet_name=sheet_name,
        total_rows=len(preview_rows),
        valid_rows_count=len(preview_rows) - errors,
        to_create_count=to_create,
        to_update_count=to_update,
        skipped_count=skipped,
        error_count=errors,
        detected_headers=headers,
        preview_rows=preview_rows,
        catalog_metrics=metrics
    )

@router.post("/confirm", response_model=ImportResult)
def confirm_excel_import(
    payload: ImportConfirmRequest,
    db: Session = Depends(get_db)
):
    created_count = 0
    updated_count = 0
    skipped_count = 0
    error_count = 0

    # Ensure default warehouse location exists
    alm_loc = db.query(Location).filter(Location.code == "ALM-01").first()
    if not alm_loc:
        alm_loc = Location(code="ALM-01", name="Almacén Principal", type="ALMACEN")
        db.add(alm_loc)
        db.flush()

    try:
        # Atomic database transaction
        with db.begin_nested():
            for row in payload.preview_rows:
                if row.action == "ERROR":
                    error_count += 1
                    continue

                if row.action == "OMITIR":
                    skipped_count += 1
                    continue

                # Category lookup or creation
                cat_id = None
                if row.category:
                    cat_name = row.category.strip()
                    category = db.query(Category).filter(Category.name.ilike(cat_name)).first()
                    if not category:
                        category = Category(name=cat_name, description=f"Categoría importada de {payload.sheet_name}")
                        db.add(category)
                        db.flush()
                    cat_id = category.id

                monthly_proj_json = None
                if row.auxiliary_data and "monthly_projections" in row.auxiliary_data:
                    monthly_proj_json = json.dumps(row.auxiliary_data["monthly_projections"])

                existing_p = db.query(Product).filter(Product.sku == row.sku.upper().strip()).first()

                if existing_p:
                    # Update product metadata if new info provided
                    if row.action in ["ACTUALIZAR", "CREAR"]:
                        if cat_id and not existing_p.category_id:
                            existing_p.category_id = cat_id
                        if monthly_proj_json and not existing_p.monthly_projections:
                            existing_p.monthly_projections = monthly_proj_json
                        updated_count += 1
                else:
                    # Create product
                    new_p = Product(
                        sku=row.sku.strip(),
                        name=row.name.strip(),
                        description=f"Importado de hoja {row.origin_sheet}",
                        category_id=cat_id,
                        unit_of_measure="unidades",
                        total_quantity=0, # Excel projections & sales DO NOT increment stock directly
                        reserved_quantity=0,
                        min_stock=5,
                        default_location_code="ALM-01",
                        annual_physical_count=0, # Excel import DOES NOT increment physical count counter
                        monthly_projections=monthly_proj_json,
                        source_origin=row.origin_sheet
                    )
                    db.add(new_p)
                    db.flush()

                    # Stock location entry
                    db.add(StockLocation(
                        product_id=new_p.id,
                        location_id=alm_loc.id,
                        quantity=0
                    ))
                    created_count += 1

                # If DATA VENTAS, store sales history entry
                if "VENTAS" in row.origin_sheet.upper() and row.auxiliary_data:
                    sale_date_val = None
                    if row.auxiliary_data.get("sale_date"):
                        try:
                            sale_date_val = datetime.fromisoformat(row.auxiliary_data["sale_date"])
                        except Exception:
                            sale_date_val = None

                    db.add(SalesHistory(
                        order_number=row.auxiliary_data.get("order_number"),
                        sale_date=sale_date_val,
                        sku=row.sku.strip(),
                        description=row.name.strip(),
                        quantity=row.auxiliary_data.get("quantity", 1),
                        company=row.auxiliary_data.get("company"),
                        city=row.auxiliary_data.get("city"),
                        category_raw=row.category
                    ))

        # Commit overall transaction
        db.commit()

        # Log import result
        import_log = ExcelImportLog(
            filename=payload.filename,
            sheet_name=payload.sheet_name,
            imported_by=payload.user_id,
            total_records=len(payload.preview_rows),
            created_count=created_count,
            updated_count=updated_count,
            skipped_count=skipped_count,
            error_count=error_count,
            details_json=json.dumps({
                "message": "Importación completada exitosamente.",
                "total_processed": len(payload.preview_rows)
            })
        )
        db.add(import_log)
        db.commit()

        metrics = calculate_catalog_metrics(db)

        return ImportResult(
            log_id=import_log.id,
            filename=payload.filename,
            sheet_name=payload.sheet_name,
            created_count=created_count,
            updated_count=updated_count,
            skipped_count=skipped_count,
            error_count=error_count,
            catalog_metrics=metrics,
            message=f"Importación completada con éxito. Creados: {created_count}, Actualizados: {updated_count}, Omitidos: {skipped_count}, Errores: {error_count}."
        )

    except Exception as e:
        db.rollback()
        raise HTTPException(status_code=500, detail=f"Fallo en la transacción de importación: {str(e)}")

@router.get("/logs", response_model=List[ExcelImportLogOut])
def get_import_logs(db: Session = Depends(get_db)):
    return db.query(ExcelImportLog).order_by(ExcelImportLog.created_at.desc()).all()
